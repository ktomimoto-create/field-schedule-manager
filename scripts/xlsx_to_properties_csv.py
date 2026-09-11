#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""FC「物件リスト」Export (xlsx) を load_properties.mjs 用の CSV に変換する。

使い方:
  python xlsx_to_properties_csv.py "C:/Users/000367/Downloads/物件リスト (4).xlsx"
  → 同フォルダに equipment_import_YYYY-MM-DD.csv を出力

xlsx 仕様（fc-ops docs/10_equipment-refresh.md と同じ）:
  1シート・12列固定: 号機/マンション名/都道府県/部屋数/BOX数/筐体数/住所/
  運用開始日/型式名/タイプ名/プログラムVer/サービス一覧
  末尾に空行＋「適用されたフィルター: …」注記行が付くので落とす。
"""
import csv
import datetime as dt
import sys
from pathlib import Path

import openpyxl

HEADERS = ['machine_number', 'property_name', 'prefectures', 'rooms', 'box_count',
           'case_count', 'address', 'start_date', 'model', 'type_name', 'program', 'service']


def main():
    if len(sys.argv) < 2:
        print('使い方: python xlsx_to_properties_csv.py <物件リスト.xlsx>')
        sys.exit(1)
    src = Path(sys.argv[1])
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = []
    skipped = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # ヘッダ行
        vals = ['' if v is None else str(v).strip() for v in (list(row) + [''] * 12)[:12]]
        machine = vals[0]
        # 末尾の空行・「適用されたフィルター」等の注記行を除外（号機が数字でない行）
        if not machine.isdigit():
            skipped += 1
            continue
        if isinstance(row[7], (dt.date, dt.datetime)):
            vals[7] = row[7].strftime('%Y-%m-%d')
        rows.append(vals)

    out = src.parent / f'equipment_import_{dt.date.today():%Y-%m-%d}.csv'
    with open(out, 'w', encoding='utf-8-sig', newline='') as fh:
        w = csv.writer(fh)
        w.writerow(HEADERS)
        w.writerows(rows)

    print(f'変換完了: {len(rows)} 件（除外 {skipped} 行） -> {out}')


if __name__ == '__main__':
    main()
